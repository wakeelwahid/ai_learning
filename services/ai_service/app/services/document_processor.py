"""
Server-side document text extraction + chunking.

Supports PDF, PPTX, DOCX, TXT (and a best-effort fallback for other text files).
Used by the upload-processing celery task so ingestion is fully self-contained —
no external worker required.
"""
import io
import re


def extract_text(data: bytes, filename: str) -> str:
    """Flat text (all pages joined)."""
    return "\n\n".join(t for t, _ in extract_pages(data, filename))


def extract_pages(data: bytes, filename: str) -> list[tuple[str, int | None]]:
    """
    Return [(text, page)] segments so chunks can cite a page/slide number.
    PDF → per page; PPTX → per slide; DOCX/TXT → single segment (page None).
    """
    ext = (filename or "").rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return _pdf_pages(data)
    if ext == "pptx":
        return _pptx_slides(data)
    if ext == "docx":
        return [(_from_docx(data), None)]
    if ext in ("txt", "md", "csv"):
        return [(data.decode("utf-8", errors="ignore"), None)]
    try:
        return [(data.decode("utf-8", errors="ignore"), None)]
    except Exception:
        return []


def _pdf_pages(data: bytes) -> list[tuple[str, int | None]]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    out = []
    for i, page in enumerate(reader.pages):
        try:
            txt = (page.extract_text() or "").strip()
        except Exception:
            txt = ""
        if txt:
            out.append((txt, i + 1))
    return out


def _pptx_slides(data: bytes) -> list[tuple[str, int | None]]:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in para.runs).strip()
                    if t:
                        parts.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        txt = "\n".join(parts).strip()
        if txt:
            out.append((txt, i + 1))
    return out


def _from_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def chunk_text(text: str, chunk_chars: int = 2200, overlap: int = 250) -> list[str]:
    """Sentence-aware greedy chunking with overlap."""
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []

    sentences = re.split(r"(?<=[.!?])\s+|\n{2,}", text)
    chunks: list[str] = []
    cur = ""
    for s in sentences:
        s = s.strip()
        if not s:
            continue
        if len(cur) + len(s) + 1 <= chunk_chars:
            cur = f"{cur} {s}".strip()
        else:
            if cur:
                chunks.append(cur)
            # start next chunk with tail overlap of the previous one
            tail = cur[-overlap:] if cur and overlap else ""
            cur = f"{tail} {s}".strip()
            # a single very long sentence: hard-split
            while len(cur) > chunk_chars:
                chunks.append(cur[:chunk_chars])
                cur = cur[chunk_chars - overlap:]
    if cur:
        chunks.append(cur)
    return [c for c in chunks if len(c) >= 30]
